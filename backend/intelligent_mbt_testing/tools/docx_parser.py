# -*- coding: utf-8 -*-
"""
DocxParser - Word文档解析工具
专门用于解析需求文档，提取结构化信息
"""

import json
import re
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple
from dataclasses import dataclass
from datetime import datetime

try:
    from docx import Document
    from docx.shared import Inches
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

@dataclass
class RequirementSection:
    """需求章节"""
    title: str
    content: str
    level: int
    subsections: List['RequirementSection']
    requirements: List[Dict[str, Any]]

@dataclass
class ParsedRequirement:
    """解析后的需求"""
    id: str
    title: str
    description: str
    priority: str
    category: str
    acceptance_criteria: List[str]
    dependencies: List[str]
    test_scenarios: List[str]

class DocxParser:
    """Word文档解析器"""
    
    def __init__(self):
        if not DOCX_AVAILABLE:
            raise ImportError("需要安装python-docx: pip install python-docx")
        
        # 需求识别模式
        self.requirement_patterns = {
            'requirement_id': r'(?:需求|REQ|R)[-_]?\d+',
            'priority': r'(?:优先级|priority)[:：]\s*(?:高|中|低|high|medium|low|critical)',
            'acceptance_criteria': r'(?:验收标准|acceptance criteria|AC)[:：]',
            'test_scenario': r'(?:测试场景|test scenario|TS)[:：]',
            'dependency': r'(?:依赖|dependency|depends on)[:：]'
        }
        
        # 章节标题模式
        self.section_patterns = [
            r'^\d+\.?\s+.*',  # 1. 章节标题
            r'^第\d+章\s+.*',  # 第1章 标题
            r'^[一二三四五六七八九十]+[、\.]\s+.*',  # 一、标题
            r'^[A-Z]+\.?\s+.*'  # A. 标题
        ]
        
        print("📄 DocxParser初始化完成")
    
    def parse_document(self, file_path: str) -> Dict[str, Any]:
        """解析Word文档"""
        try:
            if not Path(file_path).exists():
                raise FileNotFoundError(f"文件不存在: {file_path}")
            
            print(f"📖 开始解析文档: {file_path}")
            
            # 打开文档
            doc = Document(file_path)
            
            # 提取基础信息
            doc_info = self._extract_document_info(doc, file_path)
            
            # 解析文档结构
            sections = self._parse_document_structure(doc)
            
            # 提取需求
            requirements = self._extract_requirements(doc, sections)
            
            # 生成测试用例建议
            test_suggestions = self._generate_test_suggestions(requirements)
            
            # 构建解析结果
            parsed_result = {
                'document_info': doc_info,
                'sections': [self._section_to_dict(s) for s in sections],
                'requirements': [self._requirement_to_dict(r) for r in requirements],
                'test_suggestions': test_suggestions,
                'parsing_metadata': {
                    'parsed_at': datetime.now().isoformat(),
                    'total_paragraphs': len(doc.paragraphs),
                    'total_sections': len(sections),
                    'total_requirements': len(requirements),
                    'parser_version': '1.0.0'
                }
            }
            
            print(f"✅ 文档解析完成: {len(sections)}个章节, {len(requirements)}个需求")
            return parsed_result
            
        except Exception as e:
            print(f"❌ 文档解析失败: {str(e)}")
            raise
    
    def _extract_document_info(self, doc: Document, file_path: str) -> Dict[str, Any]:
        """提取文档基础信息"""
        file_path = Path(file_path)
        
        # 从文档属性提取信息
        core_props = doc.core_properties
        
        # 从文档开头提取项目信息
        project_info = self._extract_project_info(doc)
        
        return {
            'filename': file_path.name,
            'file_path': str(file_path.absolute()),
            'file_size': file_path.stat().st_size if file_path.exists() else 0,
            'title': core_props.title or project_info.get('title', file_path.stem),
            'author': core_props.author or project_info.get('author', ''),
            'subject': core_props.subject or project_info.get('subject', ''),
            'created': core_props.created.isoformat() if core_props.created else '',
            'modified': core_props.modified.isoformat() if core_props.modified else '',
            'version': project_info.get('version', '1.0'),
            'project_name': project_info.get('project_name', ''),
            'document_type': project_info.get('document_type', '需求文档')
        }
    
    def _extract_project_info(self, doc: Document) -> Dict[str, str]:
        """从文档开头提取项目信息"""
        project_info = {}
        
        # 检查前10个段落
        for i, para in enumerate(doc.paragraphs[:10]):
            text = para.text.strip()
            if not text:
                continue
                
            # 提取项目名称
            if '项目' in text and '名称' in text:
                project_info['project_name'] = re.sub(r'.*项目名称[:：]\s*', '', text)
            
            # 提取版本
            if '版本' in text:
                version_match = re.search(r'版本[:：]\s*([v\d\.]+)', text)
                if version_match:
                    project_info['version'] = version_match.group(1)
            
            # 提取文档类型
            if '需求' in text and ('文档' in text or '规格' in text):
                project_info['document_type'] = '需求文档'
            elif '设计' in text and '文档' in text:
                project_info['document_type'] = '设计文档'
        
        return project_info
    
    def _parse_document_structure(self, doc: Document) -> List[RequirementSection]:
        """解析文档结构"""
        sections = []
        current_section = None
        current_subsection = None
        
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            
            # 判断是否为标题
            title_level = self._get_title_level(para, text)
            
            if title_level > 0:
                # 创建新章节
                new_section = RequirementSection(
                    title=text,
                    content='',
                    level=title_level,
                    subsections=[],
                    requirements=[]
                )
                
                if title_level == 1:
                    # 一级标题
                    sections.append(new_section)
                    current_section = new_section
                    current_subsection = None
                elif title_level == 2 and current_section:
                    # 二级标题
                    current_section.subsections.append(new_section)
                    current_subsection = new_section
                elif title_level >= 3 and current_subsection:
                    # 三级及以下标题
                    current_subsection.subsections.append(new_section)
            else:
                # 普通内容
                if current_subsection:
                    current_subsection.content += text + '\n'
                elif current_section:
                    current_section.content += text + '\n'
        
        return sections
    
    def _get_title_level(self, para, text: str) -> int:
        """判断标题级别"""
        # 检查样式
        if para.style and 'heading' in para.style.name.lower():
            return int(para.style.name.split()[-1])
        
        # 检查文本模式
        for i, pattern in enumerate(self.section_patterns, 1):
            if re.match(pattern, text):
                return i
        
        # 检查格式（粗体、大字体等）
        if para.runs:
            first_run = para.runs[0]
            if first_run.bold and len(text) < 100:
                return 2
        
        return 0
    
    def _extract_requirements(self, doc: Document, sections: List[RequirementSection]) -> List[ParsedRequirement]:
        """提取需求"""
        requirements = []
        req_id_counter = 1
        
        for section in sections:
            section_requirements = self._extract_requirements_from_section(section, req_id_counter)
            requirements.extend(section_requirements)
            req_id_counter += len(section_requirements)
        
        return requirements
    
    def _extract_requirements_from_section(self, section: RequirementSection, start_id: int) -> List[ParsedRequirement]:
        """从章节中提取需求"""
        requirements = []
        content = section.content
        
        # 分割成需求块
        req_blocks = self._split_into_requirement_blocks(content)
        
        for i, block in enumerate(req_blocks):
            if not block.strip():
                continue
                
            # 解析需求
            req = self._parse_requirement_block(block, f"REQ_{start_id + i:03d}", section.title)
            if req:
                requirements.append(req)
        
        # 递归处理子章节
        for subsection in section.subsections:
            sub_requirements = self._extract_requirements_from_section(subsection, start_id + len(requirements))
            requirements.extend(sub_requirements)
        
        return requirements
    
    def _split_into_requirement_blocks(self, content: str) -> List[str]:
        """将内容分割为需求块"""
        # 按照需求标识符分割
        blocks = []
        lines = content.split('\n')
        current_block = []
        
        for line in lines:
            line = line.strip()
            if not line:
                continue
            
            # 检查是否为新需求的开始
            if self._is_requirement_start(line):
                if current_block:
                    blocks.append('\n'.join(current_block))
                current_block = [line]
            else:
                current_block.append(line)
        
        if current_block:
            blocks.append('\n'.join(current_block))
        
        return blocks
    
    def _is_requirement_start(self, line: str) -> bool:
        """判断是否为需求开始"""
        # 检查需求标识符
        if re.search(self.requirement_patterns['requirement_id'], line):
            return True
        
        # 检查需求关键词
        requirement_keywords = ['需求', '功能', '特性', 'requirement', 'feature', 'function']
        return any(keyword in line.lower() for keyword in requirement_keywords)
    
    def _parse_requirement_block(self, block: str, default_id: str, category: str) -> Optional[ParsedRequirement]:
        """解析需求块"""
        lines = block.split('\n')
        if not lines:
            return None
        
        # 提取需求ID
        req_id = default_id
        id_match = re.search(self.requirement_patterns['requirement_id'], block)
        if id_match:
            req_id = id_match.group(0)
        
        # 提取标题（通常是第一行）
        title = lines[0].strip()
        title = re.sub(r'^(?:需求|REQ|R)[-_]?\d+[:：]?\s*', '', title)
        
        # 提取描述
        description_lines = []
        acceptance_criteria = []
        test_scenarios = []
        dependencies = []
        priority = 'medium'
        
        current_section = 'description'
        
        for line in lines[1:]:
            line = line.strip()
            if not line:
                continue
            
            # 检查特殊标记
            if re.search(self.requirement_patterns['acceptance_criteria'], line, re.IGNORECASE):
                current_section = 'acceptance'
                continue
            elif re.search(self.requirement_patterns['test_scenario'], line, re.IGNORECASE):
                current_section = 'test'
                continue
            elif re.search(self.requirement_patterns['dependency'], line, re.IGNORECASE):
                current_section = 'dependency'
                continue
            elif re.search(self.requirement_patterns['priority'], line, re.IGNORECASE):
                priority_match = re.search(r'(?:高|中|低|high|medium|low|critical)', line.lower())
                if priority_match:
                    priority_map = {
                        '高': 'high', 'high': 'high', 'critical': 'critical',
                        '中': 'medium', 'medium': 'medium',
                        '低': 'low', 'low': 'low'
                    }
                    priority = priority_map.get(priority_match.group(0), 'medium')
                continue
            
            # 根据当前章节分类内容
            if current_section == 'description':
                description_lines.append(line)
            elif current_section == 'acceptance':
                acceptance_criteria.append(line)
            elif current_section == 'test':
                test_scenarios.append(line)
            elif current_section == 'dependency':
                dependencies.append(line)
        
        return ParsedRequirement(
            id=req_id,
            title=title,
            description='\n'.join(description_lines),
            priority=priority,
            category=category,
            acceptance_criteria=acceptance_criteria,
            dependencies=dependencies,
            test_scenarios=test_scenarios
        )
    
    def _generate_test_suggestions(self, requirements: List[ParsedRequirement]) -> List[Dict[str, Any]]:
        """生成测试用例建议"""
        test_suggestions = []
        
        for req in requirements:
            # 基于需求生成测试建议
            suggestion = {
                'requirement_id': req.id,
                'requirement_title': req.title,
                'suggested_test_cases': [],
                'test_types': [],
                'priority': req.priority,
                'estimated_effort': self._estimate_test_effort(req)
            }
            
            # 基于验收标准生成测试用例
            for i, criteria in enumerate(req.acceptance_criteria, 1):
                test_case = {
                    'id': f"{req.id}_TC_{i:02d}",
                    'name': f"验证{req.title} - {criteria[:30]}...",
                    'description': f"验证需求 {req.id} 的验收标准: {criteria}",
                    'priority': req.priority,
                    'steps': self._generate_test_steps(req, criteria),
                    'expected_result': criteria
                }
                suggestion['suggested_test_cases'].append(test_case)
            
            # 基于现有测试场景
            for i, scenario in enumerate(req.test_scenarios, len(req.acceptance_criteria) + 1):
                test_case = {
                    'id': f"{req.id}_TC_{i:02d}",
                    'name': f"测试场景 - {scenario[:30]}...",
                    'description': scenario,
                    'priority': req.priority,
                    'steps': self._generate_test_steps(req, scenario),
                    'expected_result': "按照测试场景正常执行"
                }
                suggestion['suggested_test_cases'].append(test_case)
            
            # 确定测试类型
            suggestion['test_types'] = self._determine_test_types(req)
            
            test_suggestions.append(suggestion)
        
        return test_suggestions
    
    def _estimate_test_effort(self, req: ParsedRequirement) -> str:
        """估算测试工作量"""
        complexity_score = 0
        
        # 基于描述长度
        complexity_score += min(len(req.description) // 100, 5)
        
        # 基于验收标准数量
        complexity_score += len(req.acceptance_criteria) * 2
        
        # 基于依赖关系
        complexity_score += len(req.dependencies) * 3
        
        # 基于优先级
        priority_weights = {'critical': 5, 'high': 4, 'medium': 2, 'low': 1}
        complexity_score += priority_weights.get(req.priority, 2)
        
        if complexity_score <= 5:
            return 'low'
        elif complexity_score <= 15:
            return 'medium'
        else:
            return 'high'
    
    def _generate_test_steps(self, req: ParsedRequirement, criteria: str) -> List[str]:
        """生成测试步骤"""
        steps = []
        
        # 基于需求类型生成通用步骤
        if '登录' in req.title or 'login' in req.title.lower():
            steps = [
                "打开应用程序",
                "导航到登录页面",
                "输入有效的用户名和密码",
                "点击登录按钮",
                "验证登录成功"
            ]
        elif '搜索' in req.title or 'search' in req.title.lower():
            steps = [
                "打开应用程序",
                "定位到搜索功能",
                "输入搜索关键词",
                "执行搜索",
                "验证搜索结果"
            ]
        elif '创建' in req.title or 'create' in req.title.lower():
            steps = [
                "打开应用程序",
                "导航到创建功能",
                "填写必要信息",
                "提交创建请求",
                "验证创建成功"
            ]
        else:
            # 通用步骤
            steps = [
                "准备测试环境",
                "执行测试操作",
                "验证结果",
                "清理测试数据"
            ]
        
        return steps
    
    def _determine_test_types(self, req: ParsedRequirement) -> List[str]:
        """确定测试类型"""
        test_types = ['functional']  # 默认功能测试
        
        # 基于需求内容确定测试类型
        content = (req.title + ' ' + req.description).lower()
        
        if any(keyword in content for keyword in ['性能', 'performance', '响应时间', 'speed']):
            test_types.append('performance')
        
        if any(keyword in content for keyword in ['安全', 'security', '权限', 'auth']):
            test_types.append('security')
        
        if any(keyword in content for keyword in ['界面', 'ui', '用户体验', 'ux', '交互']):
            test_types.append('ui')
        
        if any(keyword in content for keyword in ['接口', 'api', '集成', 'integration']):
            test_types.append('integration')
        
        if any(keyword in content for keyword in ['兼容', 'compatibility', '浏览器', 'browser']):
            test_types.append('compatibility')
        
        return test_types
    
    def _section_to_dict(self, section: RequirementSection) -> Dict[str, Any]:
        """将章节转换为字典"""
        return {
            'title': section.title,
            'content': section.content,
            'level': section.level,
            'subsections': [self._section_to_dict(s) for s in section.subsections],
            'requirements_count': len(section.requirements)
        }
    
    def _requirement_to_dict(self, req: ParsedRequirement) -> Dict[str, Any]:
        """将需求转换为字典"""
        return {
            'id': req.id,
            'title': req.title,
            'description': req.description,
            'priority': req.priority,
            'category': req.category,
            'acceptance_criteria': req.acceptance_criteria,
            'dependencies': req.dependencies,
            'test_scenarios': req.test_scenarios
        }

def parse_docx_file(file_path: str) -> Dict[str, Any]:
    """便捷函数：解析docx文件"""
    parser = DocxParser()
    return parser.parse_document(file_path)

if __name__ == "__main__":
    # 测试解析器
    import sys
    
    if len(sys.argv) > 1:
        file_path = sys.argv[1]
        try:
            result = parse_docx_file(file_path)
            print(json.dumps(result, ensure_ascii=False, indent=2))
        except Exception as e:
            print(f"解析失败: {e}")
    else:
        print("用法: python docx_parser.py <docx_file_path>")
